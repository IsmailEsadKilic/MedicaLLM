"""
Parse tr_drug_brand_mapping.xlsx and extract clean Turkish drug brand names.

Strategy for brand name extraction:
- Strip content in parentheses.
- Walk the tokens left-to-right and stop at the first token that is either:
    * contains a digit
    * is the "%" symbol or starts with "%"
    * is a recognized dosage-form / unit keyword (TABLET, KAPSUL, COZELTI, ...)
- The remaining prefix tokens are considered the brand name.

Usage:
    python parse_tr_brands.py [--preview N] [--out csv_path]
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl

# Tokens that indicate the start of dosage form / presentation, not part of the brand
STOP_KEYWORDS = {
    # forms
    "TABLET", "TABLETI", "TABLETLER",
    "KAPSUL", "KAPSULU", "KAPSULLER", "KAPSULUN",
    "COZELTI", "COZELTISI", "SOLUSYON", "SOLUSYONU",
    "SUSPANSIYON", "EMULSIYON", "EMULSION",
    "KREM", "POMAT", "POMADI", "MERHEM", "JEL", "GEL",
    "URUP", "SIRUP", "SYRUP",
    "DAMLA", "DAMLASI", "DAMLALARI",
    "SPREY", "AEROSOL",
    "INHALER", "INHALAT", "INHALASYON",
    "INFUZYON", "INFUZYONLUK",
    "ENJEKSIYON", "ENJEKSIYONLUK", "ENJ.",
    "LIYOFILIZE", "LIOFILIZE", "LYOPHILIZED",
    "TOZ", "TOZU", "POWDER",
    "FLAKON", "AMPUL", "AMPULU",
    "FILM", "KAPLI", "DRAJE",
    "YUMUSAK", "SERT", "EFERVESAN", "EFFERVESAN",
    "ORAL", "OFTALMIK", "NAZAL", "OTIK", "TOPIKAL", "VAJINAL", "REKTAL",
    "DERMAL", "TRANSDERMAL",
    "LOKMAN", "PASTIL", "SAKIZ",
    "SACHET", "POSET", "POSETI",
    "SUPOZITUAR", "SUPPOSITUAR", "OVUL",
    "BANT",
    "GRANUL", "GRANULE",
    "KONSANTRE",
    "LAVMAN",
    "MERHEMI",
    "SISE",
    # routes
    "I.V.", "I.M.", "S.C.", "IV", "IM", "SC",
    "INTRAVENOZ", "INTRAMUSKULER", "SUBKUTAN",
    # misc cut-words commonly appearing after dose
    "ICIN", "HAZIRLAMAK", "HAZIRLAMADA", "KULLANILACAK",
    "MODIFIYE", "SALIMLI", "UZATILMIS",
}

DIGIT_RE = re.compile(r"\d")
# Tokens like "MG", "ML", "MCG", "IU", "MG/ML"
UNIT_RE = re.compile(r"^(MG|MCG|G|KG|ML|L|IU|U|MEQ|MBQ|KIE|%)(/.+)?$", re.IGNORECASE)


def _extract_brand_from_tokens(tokens: list[str]) -> str:
    brand_tokens: list[str] = []
    for tok in tokens:
        upper = tok.upper()
        if not upper:
            continue
        if DIGIT_RE.search(upper):
            break
        if UNIT_RE.match(upper):
            break
        if upper.startswith("%"):
            break
        if upper in STOP_KEYWORDS:
            break
        brand_tokens.append(tok)
    return " ".join(brand_tokens).strip(" -,.;:")


def clean_brand(raw: str) -> str:
    if not raw:
        return ""
    name = str(raw).strip()
    # Remove content in parentheses (e.g. "(30 G)", "(1 FLAKON)")
    name = re.sub(r"\([^)]*\)", " ", name)
    # Normalize whitespace
    name = re.sub(r"\s+", " ", name).strip()

    tokens = name.split(" ")
    brand = _extract_brand_from_tokens(tokens)
    if brand:
        return brand

    # Fallback 1: the whole name may be dash-separated without spaces
    # e.g. "D-COLEFOR-10000-IU-YUMUSAK-KAPSUL"
    if "-" in name:
        dash_tokens = re.split(r"-", name)
        dash_brand = _extract_brand_from_tokens([t for t in dash_tokens if t])
        if dash_brand:
            return dash_brand.replace(" ", "-")

    # Fallback 2: name starts with digits like "5-FROTU ..." — keep leading digit-prefixed
    # token if it contains letters too
    first = tokens[0] if tokens else ""
    if first and re.search(r"[A-Za-zÇĞİÖŞÜçğıöşü]", first):
        # allow first token even if it has digits, then continue normal logic after it
        rest_brand = _extract_brand_from_tokens(tokens[1:])
        return (first + (" " + rest_brand if rest_brand else "")).strip(" -,.;:")

    # Fallback 3: pure generic like "% 0.9 SODYUM KLORUR SOLUSYONU 500 ML" — take all
    # alphabetic tokens up to the first stop/unit/digit-starting token
    alpha_tokens: list[str] = []
    for tok in tokens:
        upper = tok.upper()
        if upper.startswith("%") or not upper:
            continue
        if DIGIT_RE.search(upper) or UNIT_RE.match(upper) or upper in STOP_KEYWORDS:
            if alpha_tokens:
                break
            continue
        alpha_tokens.append(tok)
    if alpha_tokens:
        return " ".join(alpha_tokens).strip(" -,.;:")

    return ""


def normalize_atc_name(atc_name: str) -> list[str]:
    """ATC names can be like 'acetaminophen, chlorphenamine and pseudoephedrine'.
    Split on commas/'and'/'+' into separate ingredient names (lowercase)."""
    if not atc_name:
        return []
    s = str(atc_name).lower()
    # Replace separators with comma
    s = re.sub(r"\b(and|ve|\+|/)\b", ",", s)
    parts = [p.strip() for p in s.split(",") if p.strip()]
    return parts


def iter_rows(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    header_row_idx = None
    headers: list[str] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first = (row[0] or "").strip() if isinstance(row[0], str) else ""
        if first == "İlaç Adı":
            header_row_idx = i
            headers = [str(c or "").strip() for c in row]
            break
    if header_row_idx is None:
        raise RuntimeError("Header row 'İlaç Adı' not found")

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row or not row[0]:
            continue
        record = {headers[i]: row[i] for i in range(len(headers))}
        yield record


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="/Users/dogukang/Desktop/MedicaLLM-master/tr_drug_brand_mapping.xlsx")
    ap.add_argument("--preview", type=int, default=25, help="Print N sample rows")
    ap.add_argument("--out", type=str, default="", help="Write cleaned CSV to this path")
    ap.add_argument("--search", type=str, default="", help="Only show rows whose original name matches this substring")
    args = ap.parse_args()

    xlsx_path = Path(args.xlsx)
    rows = list(iter_rows(xlsx_path))
    print(f"Total rows: {len(rows)}")

    cleaned = []
    for r in rows:
        full = r.get("İlaç Adı", "") or ""
        brand = clean_brand(full)
        atc_code = r.get("ATC Kodu") or ""
        atc_name = (r.get("ATC Adı") or "").strip()
        ingredients = normalize_atc_name(atc_name)
        cleaned.append(
            {
                "brand_name": brand,
                "full_product_name": full,
                "barcode": str(r.get("Barkod") or "").strip(),
                "atc_code": str(atc_code).strip(),
                "atc_name": atc_name,
                "ingredients": ingredients,
                "company": (r.get("Firma Adı") or "").strip(),
                "prescription_type": (r.get("Reçete Türü") or "").strip(),
                "status": (r.get("Durumu") or "").strip(),
            }
        )

    # Search filter
    if args.search:
        q = args.search.lower()
        matched = [c for c in cleaned if q in c["full_product_name"].lower() or q in c["brand_name"].lower()]
        print(f"Matched rows for '{args.search}': {len(matched)}")
        for c in matched[:50]:
            print(f"  BRAND='{c['brand_name']}' | ATC={c['atc_code']} ({c['atc_name']}) | FULL='{c['full_product_name']}'")
        return

    # Preview
    print(f"\n=== Preview of first {args.preview} cleaned rows ===")
    for c in cleaned[: args.preview]:
        print(f"  '{c['brand_name']}'  <==  '{c['full_product_name'][:80]}' | ATC={c['atc_code']} ({c['atc_name']})")

    # Distinct brand stats
    distinct_brands = {c["brand_name"] for c in cleaned if c["brand_name"]}
    print(f"\nDistinct brand names: {len(distinct_brands)}")

    # Empty-brand diagnostics
    empty = [c for c in cleaned if not c["brand_name"]]
    print(f"Empty brand names: {len(empty)}")
    for c in empty[:5]:
        print(f"  EMPTY <=> '{c['full_product_name']}'")

    if args.out:
        with open(args.out, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["brand_name", "full_product_name", "barcode", "atc_code", "atc_name", "ingredients", "company", "prescription_type", "status"])
            for c in cleaned:
                w.writerow([c["brand_name"], c["full_product_name"], c["barcode"], c["atc_code"], c["atc_name"], "|".join(c["ingredients"]), c["company"], c["prescription_type"], c["status"]])
        print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
